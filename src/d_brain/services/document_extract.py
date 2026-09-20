"""Text extraction with page-level OCR and a cache beside the original."""
from __future__ import annotations

import hashlib
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path


class DocumentReadError(ValueError):
    pass


def run(args: list[str], deadline: float | None = None, timeout: int | None = None) -> str:
    if deadline is not None:
        left = deadline - time.monotonic()
        if left <= 0:
            raise TimeoutError('Превышено время чтения документа')
        timeout = left if timeout is None else min(timeout, left)
    try:
        from d_brain.services.execution import CURRENT_EXECUTION, run_bounded
        if CURRENT_EXECUTION.get() is not None:
            result = run_bounded(args, input='', cwd=Path.cwd(), env=os.environ.copy(),
                                 timeout=timeout, backend='command')
        else:
            result = subprocess.run(args, capture_output=True, text=True,
                                    timeout=timeout, check=False)
    except FileNotFoundError as e:
        raise DocumentReadError(f'Не установлена программа {args[0]}') from e
    if result.returncode:
        raise DocumentReadError(f'{args[0]}: {result.stderr.strip()[:200]}')
    return result.stdout


def extract_text(original: Path, deadline: float | None = None) -> Path:
    target = original.parent / 'текст.txt'
    digest = hashlib.sha256(original.read_bytes()).hexdigest()
    fingerprint = original.parent / 'текст.sha256'
    if (target.exists() and target.stat().st_size and fingerprint.exists()
            and fingerprint.read_text() == digest):
        return target
    suffix = original.suffix.lower()
    pages: list[str] = []
    if suffix == '.pdf':
        info = run(['pdfinfo', str(original)], deadline)
        match = re.search(r'^Pages:\s*(\d+)', info, re.M)
        if not match:
            raise DocumentReadError('Не удалось определить число страниц PDF')
        count = int(match[1])
        # Extract once. Preserve form-feed boundaries to identify scanned pages.
        raw = run(['pdftotext', '-layout', '-enc', 'UTF-8', str(original), '-'], deadline)
        native_pages = raw.split('\f')
        with tempfile.TemporaryDirectory(prefix='dbrain-pdf-') as temp:
            for index in range(count):
                text = native_pages[index].strip() if index < len(native_pages) else ''
                if len(re.sub(r'\W', '', text)) < 25:
                    prefix = str(Path(temp) / 'page')
                    run(['pdftoppm', '-f', str(index+1), '-l', str(index+1),
                         '-r', '180', '-singlefile', '-png', str(original), prefix], deadline)
                    text = run(['tesseract', prefix+'.png', 'stdout', '-l', 'rus+eng'], deadline).strip()
                    Path(prefix+'.png').unlink(missing_ok=True)
                pages.append(f'[Страница {index+1}]\n{text or "[Нет распознанного текста]"}')
    elif suffix in {'.txt', '.md', '.csv', '.json', '.log', '.conf', '.yaml', '.yml'}:
        data = original.read_bytes()
        try:
            pages = [data.decode('utf-8-sig')]
        except UnicodeDecodeError:
            pages = [data.decode('cp1251')]
    elif suffix == '.docx':
        from docx import Document
        doc = Document(original)
        pages = ['\n'.join(p.text for p in doc.paragraphs)]
        pages.extend('\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows) for table in doc.tables)
    elif suffix == '.pptx':
        from pptx import Presentation
        deck = Presentation(original)
        pages = [f'[Слайд {i}]\n'+'\n'.join(s.text for s in slide.shapes if s.has_text_frame)
                 for i, slide in enumerate(deck.slides, 1)]
    else:
        raise DocumentReadError('Этот формат пока не читается. Поддерживаются PDF, DOCX, PPTX и текстовые файлы.')
    text = '\n\n'.join(pages).strip()
    if not text or not re.search(r'[а-яА-ЯёЁa-zA-Z]{3}', re.sub(r'\[.*?\]', '', text)):
        raise DocumentReadError('Не удалось извлечь текст: проверь качество скана или защиту документа.')
    temporary = target.with_suffix('.tmp')
    temporary.write_text(text, encoding='utf-8')
    temporary.replace(target)
    fingerprint.write_text(digest)
    return target
