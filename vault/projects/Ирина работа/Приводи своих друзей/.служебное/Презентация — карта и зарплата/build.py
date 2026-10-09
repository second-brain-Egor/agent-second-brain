"""Один слайд 16:9: механизм акции «Приводи своих друзей» для карты с зарплатой."""
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(sys.argv[1])
IMG = Path(__file__).parent / "img"

DARK = RGBColor(0x2B, 0x60, 0x30)   # основной зелёный, Pantone 350C
GRASS = RGBColor(0x23, 0x84, 0x41)  # травяной, 348C
APPLE = RGBColor(0x6A, 0xA7, 0x44)  # яблочный
LIME = RGBColor(0xA6, 0xCE, 0x39)   # лаймовый
SUN = RGBColor(0xFF, 0xCB, 0x05)    # солнечный жёлтый, 116C
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x26, 0x32, 0x28)
MUTED = RGBColor(0x5E, 0x6B, 0x60)
TINT = RGBColor(0xEE, 0xF5, 0xEA)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes


def box(x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None, name=None):
    sh = S.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line[0]
        sh.line.width = Pt(line[1])
    sh.shadow.inherit = False
    for ref in sh._element.xpath("./p:style/a:effectRef"):
        ref.set("idx", "0")  # плоско, без тени темы
    if radius is not None:
        sh.adjustments[0] = radius
    if name:
        sh.name = name
    sh.text_frame.text = ""
    return sh


def text(x, y, w, h, runs, size=13, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, name=None, shape=None):
    """runs: строка, либо список абзацев; абзац — строка или список кусков (текст, {опции})."""
    tb = shape or S.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        pieces = para if isinstance(para, list) else [(para, {})]
        for piece, opt in pieces:
            r = p.add_run()
            r.text = piece
            f = r.font
            f.name = FONT
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = opt.get("color", color)
    return tb


def picture(path, x, y, w=None, h=None, name=None):
    pic = S.add_picture(str(path), Inches(x), Inches(y),
                        Inches(w) if w else None, Inches(h) if h else None)
    if name:
        pic.name = name
    return pic


# ── Шапка ──────────────────────────────────────────────────────────────
logo = Image.open(IMG / "logo_block.png")
logo_h = 0.62
logo_w = logo_h * logo.width / logo.height
picture(IMG / "logo_block.png", 0.5, 0.42, w=logo_w, name="Логотип РСХБ")

TX = 0.5 + logo_w + 0.45
text(TX, 0.28, 13.333 - TX - 0.5, 0.6, "Приводи своих друзей", size=34, bold=True,
     color=DARK, name="Заголовок")
text(TX, 0.86, 13.333 - TX - 0.5, 0.4,
     [[("Друг получает зарплату на новую карту — ", {}),
       ("вы получаете бонусные баллы", {"bold": True, "color": GRASS})]],
     size=17, color=INK, name="Подзаголовок")

# ── Четыре шага ────────────────────────────────────────────────────────
STEPS = [
    ("QR-код другу",
     "Создайте в «Мобильном банке» ссылку или QR-код и отправьте другу"),
    ("Друг берёт карту",
     "Подаёт заявку на retail.rshb.ru, в течение 90 дней получает карту «СВОЯ карта Плюс» "
     "в отделении и подключает «Мобильный банк»"),
    ("Зарплата на карту",
     "Друг получает на новую карту зарплату от 15 000 ₽"),
    ("Баллы — вам", None),
]
CARD_Y, CARD_H = 1.55, 3.68
GAP = 0.36
CARD_W = (13.333 - 1.0 - 3 * GAP) / 4
IMG_SIZE = 1.62
BADGE = 0.42

for i, (head, body) in enumerate(STEPS):
    cx = 0.5 + i * (CARD_W + GAP)
    final = body is None
    box(cx, CARD_Y, CARD_W, CARD_H, TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07,
        line=(SUN, 2.5) if final else None, name=f"Шаг {i + 1}: карточка")

    ix = cx + (CARD_W - IMG_SIZE) / 2
    iy = CARD_Y + 0.2
    picture(IMG / f"step{i + 1}.png", ix, iy, w=IMG_SIZE, h=IMG_SIZE, name=f"Шаг {i + 1}: иллюстрация")

    # номер шага и заголовок
    num = box(cx + 0.18, CARD_Y + 0.18, BADGE, BADGE, DARK, shape=MSO_SHAPE.OVAL, name=f"Шаг {i + 1}: номер")
    text(0, 0, 0, 0, str(i + 1), size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, shape=num)
    hy = iy + IMG_SIZE + 0.14
    text(cx + 0.22, hy, CARD_W - 0.44, 0.34, head, size=17, bold=True, color=DARK,
         name=f"Шаг {i + 1}: заголовок")

    by = hy + 0.46
    if not final:
        text(cx + 0.22, by, CARD_W - 0.44, CARD_Y + CARD_H - by - 0.12, body, size=12,
             color=INK, name=f"Шаг {i + 1}: текст")
    else:
        text(cx + 0.22, by - 0.12, CARD_W - 0.44, 0.52,
             [[("2 000", {"size": 30, "bold": True}), (" баллов", {"size": 15, "bold": True})]],
             color=DARK, anchor=MSO_ANCHOR.BOTTOM, name="Баллы: за друга")
        text(cx + 0.22, by + 0.42, CARD_W - 0.44, 0.24, "за каждого друга", size=12, color=INK,
             name="Баллы: пояснение")
        plaque = box(cx + 0.16, by + 0.74, CARD_W - 0.32, 0.4, SUN, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                     radius=0.5, name="Баллы: плашка")
        text(0, 0, 0, 0,
             [[("4 000", {"size": 15, "bold": True}), (" за каждого 3-го друга", {"size": 12})]],
             color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=plaque)

    # стрелка к следующему шагу
    if i < 3:
        ax = cx + CARD_W - 0.11
        ay = iy + IMG_SIZE / 2 - 0.29
        box(ax, ay, 0.58, 0.58, SUN, shape=MSO_SHAPE.OVAL, name=f"Стрелка {i + 1}")
        box(ax + 0.2, ay + 0.15, 0.2, 0.28, DARK, shape=MSO_SHAPE.CHEVRON, name=f"Стрелка {i + 1}: знак")

# ── Условия внизу ──────────────────────────────────────────────────────
INFO_Y, INFO_H = 5.52, 1.1
INFO = [
    ("Кто может стать другом",
     ["18 лет и старше", "На момент заявки нет действующих продуктов в Россельхозбанке"]),
    ("Что нужно вам",
     ["Участвовать в Программе лояльности и иметь действующий договор с Банком",
      "Оплатить картой хотя бы одну покупку за 90 дней до зарплаты друга"]),
    ("Сроки",
     ["Отправить ссылку — до 01.11.2026", "Зарплата другу — до 01.01.2027",
      "Баллы — в течение 30 дней после зарплаты"]),
]
COL_W = [3.3, 4.85, 3.8]
COL_GAP = (13.333 - 1.0 - sum(COL_W)) / 2
x = 0.5
for (head, lines), w in zip(INFO, COL_W):
    box(x, INFO_Y, 0.16, 0.16, GRASS, shape=MSO_SHAPE.OVAL, name=f"{head}: маркер")
    text(x + 0.26, INFO_Y - 0.06, w - 0.26, 0.28, head, size=13, bold=True, color=DARK,
         name=f"{head}: заголовок")
    text(x + 0.26, INFO_Y + 0.28, w - 0.26, INFO_H - 0.28,
         [[("— " + ln, {})] for ln in lines], size=11.5, color=INK, spacing=2,
         name=f"{head}: текст")
    x += w + COL_GAP

text(0.5, 6.78, 12.333, 0.36,
     "Условия для ссылок, отправленных с 02.06.2026 по 01.11.2026. Баллы начисляются за одно условие — "
     "то, которое друг выполнил первым. Номер друга считается по дню выполнения условия, после 3-го — заново. "
     "Ссылок и QR-кодов можно создать сколько угодно. Правила акции «Приводи своих друзей» в редакции с 02.10.2026.",
     size=8.5, color=MUTED, name="Сноска")

# ── Фирменный паттерн внизу ────────────────────────────────────────────
PATTERN = [(DARK, 2.4), (None, 0.12), (GRASS, 1.5), (None, 0.5), (APPLE, 0.9), (None, 0.55),
           (LIME, 1.6), (SUN, 0.4), (None, 0.45), (APPLE, 1.1), (None, 0.6), (LIME, 0.35),
           (None, 0.12), (GRASS, 0.9), (DARK, 1.85)]
scale = 13.333 / sum(w for _, w in PATTERN)
x = 0.0
for color, w in PATTERN:
    if color is not None:
        box(x, 7.36, w * scale, 0.14, color, name="Фирменный паттерн")
    x += w * scale

prs.core_properties.title = "Приводи своих друзей — карта и зарплата"
prs.core_properties.author = "Россельхозбанк"
prs.save(OUT)
print("saved", OUT)
