from __future__ import annotations

import asyncio
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from aiogram.types import Chat, Message, User

from d_brain.services.documents import DocumentStore, render_file_entry
from d_brain.services.document_extract import DocumentReadError, extract_text
from d_brain.services.document_jobs import deliver_job, process_job, user_error
from d_brain.services.document_output import DocumentOutputError, create_artifact, parse_content
from d_brain.services.session import SessionStore

REQUEST = 'Подготовить 3 слайда презентации: 1. Без резервирования. 2. С резервированием. 3. Процесс зачисления ЗП.'
CONTENT = {'slides': [
    {'title':'Выпуск и выдача без резервирования', 'bullets':['Заявление сотрудника','Оформление и выдача карты'], 'sources':'Пункт 3.2, стр. 1', 'notes':''},
    {'title':'Выпуск и выдача с резервированием', 'bullets':['Резервирование номера счёта','Оформление и выдача карты'], 'sources':'Пункт 3.2, стр. 2', 'notes':''},
    {'title':'Процесс зачисления зарплаты', 'bullets':['Передача реестра','Зачисление средств'], 'sources':'Пункт 4, стр. 3', 'notes':''}]}


def new_doc(tmp_path, name='example.txt', data=b'Example source text', request=REQUEST):
    store = DocumentStore(tmp_path)
    doc, fresh = store.receive(data, name, 7, 7, 100, request)
    assert fresh
    return store, doc


def ready_job(tmp_path):
    store, doc = new_doc(tmp_path)
    doc = store.place(doc['id'])
    job = store.enqueue(doc['id'], REQUEST, 7, 101)
    return store, doc, job


def test_requires_destination_and_deduplicates(tmp_path):
    store, doc = new_doc(tmp_path, '../../a.pdf')
    assert store.safe_path(doc['path']).is_relative_to(tmp_path/'.documents/incoming')
    with pytest.raises(ValueError):
        store.enqueue(doc['id'], 'Сделай слайды', 7, 5)
    same, fresh = store.receive(b'Example source text', 'duplicate.pdf', 7, 7, 100)
    assert not fresh and same['id'] == doc['id']
    same, fresh = store.receive(b'Example source text', 'duplicate.pdf', 7, 7, 101)
    assert fresh and same['id'] == doc['id']
    assert len(list((tmp_path/'.documents/incoming').rglob('*.pdf'))) == 1
    new_store = DocumentStore(tmp_path)
    placed = new_store.place(doc['id'], 'Работа')
    assert placed['path'].startswith('projects/Работа/Документы/')
    assert not list((tmp_path/'.documents/incoming').iterdir())
    assert new_store.place(doc['id'])['path'] == placed['path']
    other, _ = store.receive(b'Example source text', 'duplicate.pdf', 8, 8, 101)
    assert other['id'] != doc['id']


@pytest.mark.parametrize('project',['../escape','/etc','a/b','..','a\\b'])
def test_destination_traversal_rejected(tmp_path, project):
    store, doc = new_doc(tmp_path)
    with pytest.raises(ValueError):
        store.place(doc['id'], project)


def test_destination_symlink_rejected(tmp_path):
    store, doc = new_doc(tmp_path)
    (tmp_path/'projects').symlink_to('/tmp', target_is_directory=True)
    with pytest.raises(ValueError):
        store.place(doc['id'], 'escape')


def make_pdf(path, scanned=False):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    pdfmetrics.registerFont(TTFont('TestCyrillic', font))
    c=canvas.Canvas(str(path))
    c.setFont('TestCyrillic', 16)
    c.drawString(40, 760, 'Зарплатная карта. Выпуск без резервирования.')
    c.showPage()
    if scanned:
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.lib.utils import ImageReader
        img=Image.new('RGB', (1800,600),'white')
        draw=ImageDraw.Draw(img)
        draw.text((40,100), 'Процесс зачисления зарплаты', font=ImageFont.truetype(font,64), fill='black')
        draw.text((40,220), 'Реестр передаётся в банк.', font=ImageFont.truetype(font,60), fill='black')
        c.drawImage(ImageReader(img),30,450,width=540,height=180)
    else:
        c.setFont('TestCyrillic',16)
        c.drawString(40,760,'Процесс зачисления зарплаты. Передача реестра.')
    c.save()


@pytest.mark.parametrize('scanned',[False,True])
def test_real_pdf_extracts_every_page_and_caches(tmp_path, scanned, monkeypatch):
    original=tmp_path/'original.pdf'
    make_pdf(original,scanned)
    text=extract_text(original)
    data=text.read_text()
    assert '[Страница 2]' in data
    assert 'резервирования' in data
    assert 'зачисления' in data.lower()
    import d_brain.services.document_extract as module
    monkeypatch.setattr(module, 'run', Mock(side_effect=AssertionError('cache not reused')))
    assert extract_text(original)==text


def test_corrupt_pdf_fails_without_caching(tmp_path):
    original=tmp_path/'broken.pdf'; original.write_bytes(b'%PDF broken')
    with pytest.raises(DocumentReadError):
        extract_text(original)
    assert not (tmp_path/'текст.txt').exists()


def test_file_context_is_never_silent():
    assert 'source.pdf' in render_file_entry({'type':'file','path':'PDF/source.pdf','text':None})
    assert 'Сделай слайды' in render_file_entry({'path':'PDF/source.pdf','caption':'Сделай слайды'})


def test_real_artifact_and_reject_missing_topic(tmp_path):
    content=parse_content(json.dumps(CONTENT),REQUEST)
    output=create_artifact(content,REQUEST,tmp_path)
    from pptx import Presentation
    deck=Presentation(output)
    assert len(deck.slides)==3
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in deck.slides)
    with pytest.raises(DocumentOutputError):
        parse_content(json.dumps({'slides':CONTENT['slides'][:2]}),REQUEST)
    wrong=json.loads(json.dumps(CONTENT)); wrong['slides'][1]['title']='Другая тема'
    with pytest.raises(DocumentOutputError):
        parse_content(json.dumps(wrong),REQUEST)


@pytest.mark.parametrize('fmt', ['docx','pdf','md','txt'])
def test_output_formats(tmp_path, fmt):
    request='Сделай ответ в '+fmt
    content=CONTENT if fmt=='pdf' else {'title':'Ответ','text':'Текст по исходному документу'}
    assert create_artifact(content,request,tmp_path).suffix=='.'+fmt


def test_job_retry_cache_and_persisted_state(tmp_path):
    store,doc,job=ready_job(tmp_path)
    processor=SimpleNamespace(_run_backend_exec=Mock(side_effect=[
        subprocess.TimeoutExpired('model',900),json.dumps(CONTENT)]))
    process_job(store,job,processor)
    done=DocumentStore(tmp_path).job(job['id'])
    assert done['state']=='ready' and done['attempts']==2
    assert processor._run_backend_exec.call_count==2
    assert all(call.kwargs['timeout_sec'] is None for call in processor._run_backend_exec.call_args_list)
    assert store.safe_path(done['artifact']).exists()
    assert store.enqueue(doc['id'],REQUEST,7,101)['id']==job['id']
    text_entries=SessionStore(tmp_path).get_recent(7)
    assert all(e.get('text') is None for e in text_entries if e['type']=='file')


def test_permanent_error_not_retried(tmp_path):
    store,doc,job=ready_job(tmp_path)
    processor=SimpleNamespace(_run_backend_exec=Mock(side_effect=RuntimeError('unauthorized')))
    process_job(store,job,processor)
    assert store.job(job['id'])['state']=='failed'
    assert processor._run_backend_exec.call_count==1


async def test_delivery_once_and_restart_uncertainty(tmp_path):
    store,doc,job=ready_job(tmp_path)
    process_job(store,job,SimpleNamespace(_run_backend_exec=Mock(return_value=json.dumps(CONTENT))))
    bot=SimpleNamespace(send_document=AsyncMock(return_value=SimpleNamespace(message_id=901)))
    await asyncio.gather(deliver_job(store,store.job(job['id']),bot),deliver_job(store,store.job(job['id']),bot))
    assert bot.send_document.await_count==1
    assert store.job(job['id'])['state']=='sent'
    store.set_job(job['id'],state='sending')
    store.recover()
    assert store.job(job['id'])['state']=='send_unknown'
    await deliver_job(store,store.job(job['id']),bot)
    assert bot.send_document.await_count==1


async def test_send_timeout_keeps_artifact_without_retry(tmp_path):
    store,doc,job=ready_job(tmp_path)
    process_job(store,job,SimpleNamespace(_run_backend_exec=Mock(return_value=json.dumps(CONTENT))))
    bot=SimpleNamespace(send_document=AsyncMock(side_effect=TimeoutError()))
    await deliver_job(store,store.job(job['id']),bot)
    assert store.job(job['id'])['state']=='send_unknown'
    assert store.safe_path(store.job(job['id'])['artifact']).exists()
    assert bot.send_document.await_count==1


def test_restart_does_not_resume_generation_automatically(tmp_path):
    store,doc,job=ready_job(tmp_path)
    store.set_job(job['id'],state='generating',attempts=1)
    store.recover()
    assert store.job(job['id'])['state']=='failed'
    store.set_job(job['id'],state='generating',attempts=2)
    store.recover()
    assert store.job(job['id'])['state']=='failed'


def test_document_content_excluded_from_memory(tmp_path):
    from d_brain.services.memory_rag import _iter_markdown_files
    store,doc=new_doc(tmp_path,'uploaded.md',b'SECRET WHOLE DOCUMENT')
    doc=store.place(doc['id'],'Work')
    (tmp_path/'memory').mkdir()
    (tmp_path/'memory'/'note.md').write_text('description and path only')
    assert _iter_markdown_files(tmp_path)==[tmp_path/'memory'/'note.md']


async def test_text_and_voice_requests_share_persistent_destination(tmp_path,monkeypatch):
    import d_brain.bot.handlers.document as module
    settings=SimpleNamespace(vault_path=tmp_path)
    monkeypatch.setattr(module,'get_settings',lambda:settings)
    answers=AsyncMock()
    monkeypatch.setattr(Message,'answer',answers)
    message=Message(message_id=201,date=datetime.now(timezone.utc),chat=Chat(id=7,type='private'),
                    from_user=User(id=7,is_bot=False,first_name='Ирина'),text='В проект Банк')
    store,doc=new_doc(tmp_path,request=REQUEST)
    assert await module.route_document_request(message,'В проект Банк')
    placed=store.get(doc['id'])
    assert placed['path'].startswith('projects/Банк/Документы/')
    assert not store.jobs(('queued',))
    assert await module.route_document_request(message.model_copy(update={'message_id':203}), 'PowerPoint')
    assert len(store.jobs(('queued',)))==1
    # Simulate the same route used for a transcribed voice message after a restart.
    voice=message.model_copy(update={'message_id':202,'text':None})
    assert await module.route_document_request(voice,'Подготовь три слайда по этому документу')
    assert len(store.jobs(('queued',)))==1
    assert await module.route_document_request(voice.model_copy(update={'message_id':204}), 'PDF')
    assert len(store.jobs(('queued',)))==2
    # An unrelated command must not be captured by a previous document.
    SessionStore(tmp_path).append(7, 'voice', text='PDF', msg_id=204, chat_id=7)
    SessionStore(tmp_path).append(7, 'text', text='Сделай список покупок', msg_id=205, chat_id=7)
    assert not await module.route_document_request(voice,'Сделай список покупок')


async def test_caption_download_and_destination_survive_restart(tmp_path,monkeypatch):
    import io
    import d_brain.bot.handlers.document as module
    from aiogram.types import Document
    settings=SimpleNamespace(vault_path=tmp_path,treat_all_group_chats_as_work=True,work_chat_ids=[])
    monkeypatch.setattr(module,'get_settings',lambda:settings)
    answers=AsyncMock(); monkeypatch.setattr(Message,'answer',answers)
    msg=Message(message_id=111,date=datetime.now(timezone.utc),chat=Chat(id=7,type='private'),
        from_user=User(id=7,is_bot=False,first_name='Ирина'), caption=REQUEST,
        document=Document(file_id='f',file_unique_id='u',file_name='salary.pdf'))
    bot=SimpleNamespace(get_file=AsyncMock(return_value=SimpleNamespace(file_path='file')),
                        download_file=AsyncMock(return_value=io.BytesIO(b'%PDF content')))
    await module.handle_document(msg,bot)
    store=DocumentStore(tmp_path); doc=store.for_scope(7)[0]
    assert doc['state']=='destination' and doc['instructions']==REQUEST
    assert not store.jobs(('queued',))
    assert 'Куда сохранить' in answers.call_args.args[0]
    answer=msg.model_copy(update={'message_id':112,'text':'В PDF','caption':None,'document':None})
    assert await module.route_document_request(answer,'В PDF')
    assert store.get(doc['id'])['path'].startswith('PDF/')
    assert not store.jobs(('queued',))
    assert 'В каком формате' in answers.call_args.args[0]
    assert await module.route_document_request(answer.model_copy(update={'message_id':113}), 'PowerPoint')
    assert len(store.jobs(('queued',)))==1


def test_original_legacy_import_no_extra_copy(tmp_path):
    folder=tmp_path/'attachments';folder.mkdir()
    original=folder/'old.pdf';original.write_bytes(b'legacy')
    store=DocumentStore(tmp_path)
    doc=store.import_existing(original,7,7,55)
    assert store.safe_path(doc['path'])==original
    assert not list((tmp_path/'.documents/incoming').rglob('*.pdf'))
    placed=store.place(doc['id'])
    assert store.safe_path(placed['path']).read_bytes()==b'legacy'


def test_error_messages_do_not_guess_backend_or_cause():
    text=user_error(subprocess.TimeoutExpired('codex',900))
    assert 'Sonnet' not in text and 'Claude Max' not in text and 'Превышено время' in text


def test_delivery_timeout_reports_delivery_not_generation():
    assert 'Файл готов' in user_error(TimeoutError(), 'delivery')


def test_chat_error_does_not_claim_document_was_saved():
    assert 'Документ' not in user_error(TimeoutError())
    assert 'Документ и задание сохранены' in user_error(TimeoutError(), 'document')


def test_document_context_survives_day_boundary(tmp_path):
    from d_brain.services.processor import AgentProcessor
    store, doc = new_doc(tmp_path)
    processor = object.__new__(AgentProcessor)
    processor.vault_path = tmp_path
    context = processor._get_session_context(7)
    assert doc['name'] in context
    assert 'ожидает выбора папки' in context
    assert doc['name'] not in processor._get_session_context(8)


@pytest.mark.parametrize('state', ['destination', 'project'])
async def test_unrelated_request_keeps_pending_document_assignment(tmp_path, monkeypatch, state):
    import d_brain.bot.handlers.document as module
    monkeypatch.setattr(module, 'get_settings', lambda: SimpleNamespace(vault_path=tmp_path))
    monkeypatch.setattr(Message, 'answer', AsyncMock())
    store, doc = new_doc(tmp_path)
    store.update(doc['id'], state=state)
    message = Message(message_id=201, date=datetime.now(timezone.utc),
        chat=Chat(id=7,type='private'), from_user=User(id=7,is_bot=False,first_name='Ирина'),
        text='Сделай список покупок')
    assert not await module.route_document_request(message, message.text)
    assert store.get(doc['id'])['instructions'] == REQUEST
    assert store.get(doc['id'])['state'] == state


async def test_project_name_reply_to_bot_question(tmp_path, monkeypatch):
    import d_brain.bot.handlers.document as module
    monkeypatch.setattr(module, 'get_settings', lambda: SimpleNamespace(vault_path=tmp_path))
    monkeypatch.setattr(Message, 'answer', AsyncMock())
    store, doc = new_doc(tmp_path)
    store.update(doc['id'], state='project')
    reply = Message(message_id=200, date=datetime.now(timezone.utc), chat=Chat(id=7,type='private'),
        from_user=User(id=900,is_bot=True,first_name='Бот'),
        text='В какой проект сохранить? Напиши название проекта.')
    message = Message(message_id=201, date=datetime.now(timezone.utc),chat=Chat(id=7,type='private'),
        from_user=User(id=7,is_bot=False,first_name='Ирина'),text='Порядок выпуска ЗП карт',reply_to_message=reply)
    assert await module.route_document_request(message, message.text)
    assert store.get(doc['id'])['path'].startswith('projects/Порядок выпуска ЗП карт/Документы/')


@pytest.mark.parametrize('state', ['destination', 'project'])
async def test_delivery_followup_preserves_original_topics(tmp_path, monkeypatch, state):
    import d_brain.bot.handlers.document as module
    monkeypatch.setattr(module, 'get_settings', lambda: SimpleNamespace(vault_path=tmp_path))
    monkeypatch.setattr(Message, 'answer', AsyncMock())
    store, doc = new_doc(tmp_path)
    store.update(doc['id'], state=state)
    message = Message(message_id=201, date=datetime.now(timezone.utc),
        chat=Chat(id=7,type='private'), from_user=User(id=7,is_bot=False,first_name='Ирина'),
        text='Пришли мне сюда три слайда')
    assert await module.route_document_request(message, message.text)
    assert store.get(doc['id'])['instructions'] == REQUEST
    assert store.get(doc['id'])['state'] == state
    choice = message.model_copy(update={'message_id':202,
        'text':'Название проекта для документа Порядок выпуска ЗП карт'})
    assert await module.route_document_request(choice, choice.text)
    assert store.get(doc['id'])['path'].startswith('projects/Порядок выпуска ЗП карт/Документы/')
    assert not store.latest_job(doc['id'])
    assert await module.route_document_request(choice.model_copy(update={'message_id':203}), 'PowerPoint')
    assert store.latest_job(doc['id'])['request'].startswith(REQUEST)


@pytest.mark.parametrize('state', ['queued', 'generating', 'ready', 'sent'])
async def test_delivery_followup_does_not_generate_duplicate(tmp_path, monkeypatch, state):
    import d_brain.bot.handlers.document as module
    monkeypatch.setattr(module, 'get_settings', lambda: SimpleNamespace(vault_path=tmp_path))
    monkeypatch.setattr(Message, 'answer', AsyncMock())
    store, doc, job = ready_job(tmp_path)
    store.set_job(job['id'], state=state)
    message = Message(message_id=202, date=datetime.now(timezone.utc),
        chat=Chat(id=7,type='private'), from_user=User(id=7,is_bot=False,first_name='Ирина'),
        text='Пришли мне файл')
    assert await module.route_document_request(message, message.text)
    assert len(store.jobs((state,))) == 1
    assert store.latest_job(doc['id'])['id'] == job['id']
